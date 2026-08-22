from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = '/home/ubuntu/inventory-sync-prototype/docs/learning-and-blocker-journal.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Learning & Blocker Journal')
r.bold = True
r.font.size = Pt(20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Maama Junior School inventory synchronization\nAuthor: Manus AI | 22 August 2026').italic = True

def heading(text, level=1):
    doc.add_heading(text, level=level)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph()

heading('Context')
doc.add_paragraph('Parents reported receiving a mobile-money or bank notification but no confirmation from the school system. The delivery objective is to keep the support tool’s “Is the balance cleared?” answer accurate after an upstream balance or inventory change.')
heading('Learning log')
table(['Entry', 'Learning', 'Design consequence'], [
    ['1', 'A five-minute poll can leave support staff looking at stale state and creates an avoidable recurring dependency.', 'Preserve the cache and query contract but remove the polling trigger.'],
    ['2', 'An event receiver must authenticate the raw message before parsing or applying it.', 'Verify HMAC over timestamp.raw_body and use a bounded timestamp window.'],
    ['3', 'Providers retry deliveries, so a successful response must not imply that every delivery is new.', 'Deduplicate by eventId and return a successful duplicate response.'],
    ['4', '“Not cleared” and “unknown” are different support outcomes.', 'Return 404 for unknown records and balanceCleared: false only for known insufficient state.'],
])
heading('Blocker log')
table(['Blocker', 'Impact', 'Resolution or owner action'], [
    ['The upstream warehouse/payment provider was not named in the brief.', 'Provider-specific webhook registration and signature headers cannot be completed from the brief alone.', 'Use a provider-neutral adapter contract; the integration owner must map the provider callback before production.'],
    ['No production database or deployment credentials were supplied.', 'A process restart loses the prototype cache.', 'Keep the runnable prototype in memory and record durable storage as a release hardening item.'],
    ['The client deadline prohibits extensions.', 'Polling cannot be retained as a fallback.', 'Delete polling code and make webhook delivery a release prerequisite.'],
])
heading('Delivered evidence')
doc.add_paragraph('The repository contains a webhook handler, HMAC verification, replay protection, idempotent cache update, query endpoint, automated tests, an event-driven specification, a Scope Delta Analysis, and this journal. A separate confidential Adaptability Index records the pivot response.')
heading('Follow-up checklist')
doc.add_paragraph('Before production, configure the provider callback, place the service behind HTTPS, use a secret manager, persist event and current-state data, add delivery-lag and rejection metrics, and test provider retries and out-of-order events. These items do not reopen polling; they harden the chosen event-driven design.')

doc.save(out)
print(out)
